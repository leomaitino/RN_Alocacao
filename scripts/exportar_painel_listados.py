#!/usr/bin/env python3
"""
Exporta o painel de Fundos Listados (data/listados.json) para Excel no layout
da planilha 'Painel Fundos Listados' (Dados cadastrais · Preço · Performance ·
Dividend Yield · Volume · Informações do Fundo · Último provento · Indexador ·
Ativos · TIR · Status), mais abas de Stress, Carteira (FI-Infra) e Fontes.

Uso:
  python scripts/exportar_painel_listados.py                       # → output/Painel Fundos Listados - preenchido.xlsx
  python scripts/exportar_painel_listados.py --out caminho.xlsx
  python scripts/exportar_painel_listados.py --todos               # inclui candidatos não sugeridos
"""
import argparse
import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent
DATA = BASE / 'data'


def idx_txt(r):
    ix = r.get('indexadores_carteira') or {}
    return ' · '.join(f"{k} {v*100:.0f}%" for k, v in sorted(ix.items(), key=lambda kv: -kv[1])) if ix else ''


def cupom_txt(r):
    c = r.get('cupom_medio') or {}
    return ' · '.join(f"{k}{'' if k.endswith('+') else '+'}{v:.2f}%" for k, v in c.items()) if c else ''


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default=str(BASE / 'output' / 'Painel Fundos Listados - preenchido.xlsx'))
    ap.add_argument('--todos', action='store_true')
    a = ap.parse_args()

    fundos = json.loads((DATA / 'listados.json').read_text(encoding='utf-8'))
    cart = json.loads((DATA / 'listados_carteiras.json').read_text(encoding='utf-8'))
    meta = json.loads((DATA / 'meta_listados.json').read_text(encoding='utf-8'))
    sel = [f for f in fundos if f.get('no_painel') or f.get('status') == 'Sugerido' or a.todos]
    ordem = {'Aprovado': 0, 'Não Aprovado': 1, 'Sugerido': 2, 'Candidato': 3}
    sel.sort(key=lambda f: (ordem.get(f['status'], 9), f['tipo'], f['ticker']))

    rows = []
    for f in sel:
        c = cart.get(f['ticker'], {})
        rows.append({
            'Código': f['ticker'], 'Tipo': f['tipo'], 'Categoria': f['categoria'], 'Indexador': f['indexador'],
            'Nome': f['nome'], 'Característica': f"{f['tipo']} · {f['indexador']} · " + ('debêntures incentivadas' if f['tipo'] in ('FI-Infra', 'FIP-IE') else ('CRA' if f['tipo'] == 'Fiagro' else 'CRI (papel)')),
            'Administrador': f.get('administrador'), 'Gestor': f.get('gestor'), 'CNPJ': f['cnpj'], 'Tx de Adm (%)': f.get('taxa_adm'),
            'R$/Cota': f.get('preco'), 'Data preço': f.get('data_preco'),
            'Perf. 1 mês': f.get('var_1m'), 'Perf. no ano': f.get('var_ano'), 'Perf. 12 meses': f.get('var_12m'), 'Ret. total 12M (c/ proventos)': f.get('ret_total_12m'),
            'DY 1 mês': f.get('dy_1m'), 'DY 12 meses': f.get('dy_12m'), 'Nº proventos 12M': f.get('n_proventos_12m'),
            'Vol. média diária 3M (R$ mil)': f.get('volume_medio_3m_mil'), 'Peso no IFIX': None,
            'VM (R$ mil)': (f['vm'] * 1e3 if f.get('vm') is not None else None), 'PL (R$ mil)': (f['pl'] * 1e3 if f.get('pl') is not None else None), 'VM/PL': f.get('vm_pl'),
            'VP/cota': f.get('vp_cota'), 'Data VP': f.get('vp_data'), 'P/VP': f.get('p_vp'), 'P/VP mín 12M': f.get('p_vp_min_12m'), 'P/VP máx 12M': f.get('p_vp_max_12m'),
            'Último provento (R$/cota)': f.get('provento_ultimo'), 'Data últ. provento': f.get('provento_ultimo_data'), 'Proventos 12M (R$/cota)': f.get('proventos_12m'), 'Yield anualizado': f.get('yield_anualizado'),
            'Indexador carteira': idx_txt(f), 'Cupom médio emissão (SND)': cupom_txt(f), 'Prazo médio (anos)': c.get('prazo_medio_anos'), '% Lei 12.431': c.get('pct_incentivadas'),
            'Qtde de ativos': f.get('n_ativos'), 'Nº emissores': c.get('n_emissores'), 'Carteira confidencial (%)': c.get('confidencial_pct'),
            'TIR / taxa (relatório)': f.get('tir_relatorio'), 'Cotistas': f.get('cotistas'), 'Cotas emitidas': f.get('cotas_emitidas'),
            'Stress (nível)': {0: '', 1: 'baixo', 2: 'médio', 3: 'alto'}.get(f.get('stress_nivel') or 0, ''), 'Nº alertas': f.get('stress_n_alertas'), 'Exposição stress (% PL)': f.get('stress_exposicao_pct_pl'),
            'Alertas': ' | '.join(x['titulo'] for x in (f.get('stress_alertas') or [])), 'Status': f['status'], 'Fonte VP': f.get('rent_fonte_vp'), 'Fonte carteira': f.get('carteira_fonte'), 'Tese / obs': f.get('tese') or f.get('obs'),
        })
    df = pd.DataFrame(rows)

    stress = []
    for f in sel:
        for x in cart.get(f['ticker'], {}).get('alertas', []):
            stress.append({'Código': f['ticker'], 'Tipo': f['tipo'], 'Status': f['status'], 'Severidade': x['severidade'], 'Evento': x['titulo'], 'Tipo de evento': x.get('tipo'),
                           'Data': x.get('data'), 'Exposição (% PL)': x.get('exposicao_pct_pl'), 'Debêntures': ', '.join(x.get('codigos') or []), 'Descrição': x.get('descricao'), 'Fonte': x.get('fonte')})
    df_stress = pd.DataFrame(stress)

    deb = []
    for f in sel:
        for d in cart.get(f['ticker'], {}).get('ativos', []):
            deb.append({'Código fundo': f['ticker'], 'Debênture': d['codigo'], 'ISIN': d.get('isin'), 'Emissor': d.get('emissor'), 'CNPJ emissor': d.get('cnpj_emissor'), 'Indexador': d.get('indexador'),
                        'Cupom emissão (%)': d.get('taxa'), 'Vencimento': d.get('vencimento'), '% PL': d.get('pct_pl'), 'R$ mil': d.get('valor_mil'), 'Lei 12.431': d.get('incentivada'), 'Garantia': d.get('garantia'), 'Situação SND': d.get('situacao_snd'), 'Via': d.get('via')})
    df_deb = pd.DataFrame(deb)

    dr = meta.get('datas_referencia', {})
    df_fontes = pd.DataFrame([
        ['Gerado em', datetime.now().strftime('%d/%m/%Y %H:%M')],
        ['Preço / volume / proventos', f"B3 via Yahoo Finance (TICKER.SA) — até {dr.get('yahoo')}"],
        ['VP, PL, cotas (FI-Infra)', f"CVM informe diário (fundo de cotas listado) — até {dr.get('cvm_diario')}"],
        ['VP, PL, cotas (FII)', f"CVM informe mensal FII — ref. {dr.get('cvm_fii')}"],
        ['VP, PL, cotas (Fiagro)', f"CVM informe mensal FIAGRO — ref. {dr.get('cvm_fiagro')}"],
        ['Carteira (FI-Infra)', f"CVM CDA {dr.get('cda')} com look-through FIC → master; debêntures identificadas no SND/ANBIMA"],
        ['Peso no IFIX', 'não disponível (API da B3 fora do ar em out/2026) — coluna em branco'],
        ['TIR / taxa', 'taxa do relatório gerencial informada no painel original (campo manual); cupom médio de emissão vem do SND e NÃO é marcação a mercado'],
        ['Eventos de stress', 'data/listados_eventos_stress.json (curadoria manual + cruzamento com carteira)'],
        ['Regra VM/PL', 'VM = preço × cotas emitidas (CVM); VM/PL ≈ P/VP'],
    ], columns=['Item', 'Fonte / regra'])

    out = Path(a.out); out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine='openpyxl') as w:
        df.to_excel(w, sheet_name='Painel Fundos Listados', index=False)
        (df_stress if len(df_stress) else pd.DataFrame([{'info': 'sem alertas'}])).to_excel(w, sheet_name='Stress', index=False)
        (df_deb if len(df_deb) else pd.DataFrame([{'info': 'sem carteira CDA'}])).to_excel(w, sheet_name='Carteira FI-Infra', index=False)
        df_fontes.to_excel(w, sheet_name='Fontes', index=False)
        pct_cols = {'Tx de Adm (%)', 'Perf. 1 mês', 'Perf. no ano', 'Perf. 12 meses', 'Ret. total 12M (c/ proventos)', 'DY 1 mês', 'DY 12 meses', 'VM/PL', 'Yield anualizado', '% Lei 12.431', 'Carteira confidencial (%)', 'Exposição stress (% PL)', 'Exposição (% PL)', '% PL'}
        for name, frame in [('Painel Fundos Listados', df), ('Stress', df_stress), ('Carteira FI-Infra', df_deb), ('Fontes', df_fontes)]:
            ws = w.sheets[name]
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F1F1F'); cell.alignment = Alignment(vertical='center', wrap_text=True)
            ws.freeze_panes = 'B2' if name == 'Painel Fundos Listados' else 'A2'
            for j, col in enumerate(frame.columns, start=1):
                letter = get_column_letter(j)
                ws.column_dimensions[letter].width = min(48, max(10, int(frame[col].astype(str).str.len().quantile(0.9) if len(frame) else 10) + 2))
                if col in pct_cols:
                    for cell in ws[letter][1:]:
                        cell.number_format = '0.00%'
                elif col in ('VM (R$ mil)', 'PL (R$ mil)', 'Vol. média diária 3M (R$ mil)', 'R$ mil', 'Cotistas', 'Cotas emitidas'):
                    for cell in ws[letter][1:]:
                        cell.number_format = '#,##0'
                elif col in ('P/VP', 'P/VP mín 12M', 'P/VP máx 12M', 'VP/cota', 'R$/Cota', 'Último provento (R$/cota)', 'Proventos 12M (R$/cota)'):
                    for cell in ws[letter][1:]:
                        cell.number_format = '0.000'
            ws.auto_filter.ref = ws.dimensions
    print(f"✓ {out} — {len(df)} fundos | {len(df_stress)} alertas | {len(df_deb)} debêntures")


if __name__ == '__main__':
    main()
